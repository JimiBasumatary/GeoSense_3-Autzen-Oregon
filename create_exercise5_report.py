from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
DOCS = ROOT / "docs" / "exercise5"
FIGURES = DOCS / "figures"
REPORTS = ROOT / "outputs" / "reports"
OUTPUT = DOCS / "GeoSense_Phase_III_Exercise_5_Report.docx"

def add_page_number(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r._r.append(begin); r._r.append(instr); r._r.append(end)

def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")

def figure(doc, filename, caption):
    path = FIGURES / filename
    if path.exists():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(6.15))
        p = doc.add_paragraph(caption, style="Caption"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True
    else:
        doc.add_paragraph(f"[Figure missing: {path}]")

def main():
    DOCS.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.7)
        sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)
        add_page_number(sec.footer.paragraphs[0])
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("\n\nGEOSENSE AGENT\n").bold = True
    p.add_run("PHASE III â€” EXERCISE 5\n").bold = True
    p.add_run("LiDAR Terrain Analysis and Unified Site Profiling\n\n")
    p.add_run("Study area: Autzen Stadium / Eugene, Oregon, USA\n")
    p.add_run(f"Report generated: {date.today():%d %B %Y}\n\nTechnical Practical Report")
    doc.add_page_break()

    doc.add_heading("Abstract", 1)
    doc.add_paragraph(
        "This report documents a Python-based LiDAR and terrain-analysis workflow for the Autzen "
        "Stadium area of Eugene, Oregon. It covers classified LiDAR inspection, ground-point extraction, "
        "DEM generation, hillshade/slope/aspect derivation, terrain metrics, existing street-scene "
        "segmentation summaries, unified site profiles and automated testing. Results are exploratory. "
        "Street-scene predictions have not been validated against manually labelled reference data, and "
        "the imagery has not been spatially matched to the LiDAR. This report does not claim validated "
        "walkability, site suitability, canopy cover, impervious-surface area or building-height estimates."
    )

    doc.add_heading("1. Aim and Objectives", 1)
    doc.add_paragraph(
        "Aim: build and test a reproducible Python workflow combining LiDAR-derived terrain information "
        "and existing street-scene segmentation summaries in a unified site-profile reporting structure."
    )
    for x in [
        "Inspect original classified COPC/LAZ and extracted ground-point cloud.",
        "Create a DEM from ground points and derive hillshade, slope and aspect.",
        "Summarize raster metadata and terrain metrics in CSV/JSON outputs.",
        "Include existing street-scene segmentation outputs as exploratory context.",
        "Generate unified site profiles and run automated tests.",
    ]: bullet(doc, x)

    doc.add_heading("2. Study Area, Data and Software", 1)
    doc.add_paragraph(
        "The study area is around Autzen Stadium in Eugene, Oregon, USA. The LiDAR data use Oregon "
        "GIC Lambert horizontal coordinates (EPSG:2992) and NAVD88 height (EPSG:6360), with elevation "
        "values expressed in US survey feet. These values must not be interpreted as metres."
    )
    doc.add_heading("2.1 Input data", 2)
    for x in [
        "Original classified LiDAR: data/lidar/raw/autzen-classified.copc.laz.",
        "Ground LiDAR: data/lidar/processed/autzen_ground.laz.",
        "DEM: data/lidar/processed/autzen_dem.tif.",
        "Terrain derivatives: autzen_hillshade.tif, autzen_slope.tif and autzen_aspect.tif.",
        "Five Mapillary street-scene images and existing segmentation summaries, masks and overlays.",
        "Profile and analysis outputs under outputs/reports/.",
    ]: bullet(doc, x)
    doc.add_heading("2.2 Software", 2)
    doc.add_paragraph(
        "Python 3.11, PDAL, GDAL utilities, Rasterio, NumPy, JSON/CSV modules, python-docx and pytest. "
        "PyTorch inference in the current environment is CPU-based; CUDA was unavailable."
    )

    doc.add_heading("3. Methodology", 1)
    for x in [
        "Inspect original classified LiDAR point count, bounds and coordinate information.",
        "Extract ground-class points to a separate LAZ file using the PDAL pipeline.",
        "Create a DEM from ground points with nominal 3 ft by 3 ft pixels.",
        "Generate hillshade, slope and aspect rasters from the DEM using GDAL utilities.",
        "Inspect raster dimensions, CRS, NoData and summary statistics.",
        "Read existing street-scene segmentation summaries and overlays as contextual outputs.",
        "Generate terrain_metrics.csv and unified/enhanced site-profile JSON outputs.",
        "Run automated tests for the analyzer, bounding boxes, coordinate utilities, profiles and terrain indicators.",
    ]:
        doc.add_paragraph(x, style="List Number")

    doc.add_heading("4. Results", 1)
    doc.add_heading("4.1 LiDAR point clouds", 2)
    doc.add_paragraph(
        "The original classified point cloud contains 10,653,336 points. The extracted ground cloud "
        "contains 6,909,405 points, with elevations approximately 406.56â€“482.58 US survey feet. "
        "These counts describe the datasets and do not independently establish classification accuracy."
    )
    figure(doc, "Fig_01_original_lidar_metadata.png", "Figure 1. Summary of the original classified LiDAR point cloud.")
    figure(doc, "Fig_02_ground_lidar_metadata.png", "Figure 2. Summary of the extracted ground LiDAR point cloud.")

    doc.add_heading("4.2 DEM and terrain derivatives", 2)
    doc.add_paragraph(
        "The DEM is 1,143 columns by 1,553 rows with nominal 3 ft by 3 ft pixels. The latest GDAL "
        "metadata summary reports minimum elevation 406.9347 ft, maximum 482.19 ft, mean 423.3060 ft, "
        "standard deviation 7.2390 ft and 79.09% valid pixels. These statistics depend on the raster "
        "and its valid-data mask."
    )
    figure(doc, "Fig_03_dem_metadata.png", "Figure 3. DEM metadata and summary statistics.")
    doc.add_paragraph(
        "The mean slope is approximately 4.11 degrees. The arithmetic mean aspect is approximately "
        "182.46 degrees; aspect is circular data, so its arithmetic mean is not necessarily a robust "
        "dominant direction."
    )
    figure(doc, "Fig_04_terrain_derivatives.png", "Figure 4. DEM-derived hillshade, slope and aspect products.")

    doc.add_heading("4.3 Street-scene segmentation", 2)
    doc.add_paragraph(
        "Five street-scene images and associated segmentation masks/overlays are available. They are "
        "exploratory context only: no manually labelled reference dataset or accuracy assessment is "
        "available, and the images are not spatially matched to the LiDAR."
    )
    figure(doc, "Fig_05_segmentation_contact_sheet.jpg", "Figure 5. Contact sheet of existing street-scene segmentation outputs.")

    doc.add_heading("4.4 Analyzer, profiles and terrain metrics", 2)
    doc.add_paragraph(
        "Two 300 ft by 300 ft site-profile boxes were generated within the same general Autzen study area. "
        "They provide descriptive comparisons and are not proof of validated contrasting land-use settings."
    )
    figure(doc, "Fig_06_lidar_analyzer_evidence.png", "Figure 6. Evidence summary for the LiDAR analyzer.")
    figure(doc, "Fig_07_site_profile_comparison.png", "Figure 7. Descriptive comparison of the two Autzen-area profiles.")
    figure(doc, "Fig_08_unified_site_profile.png", "Figure 8. Summary of the unified site-profile JSON.")
    figure(doc, "Fig_09_terrain_metrics.png", "Figure 9. Preview of the terrain metrics CSV.")

    doc.add_heading("4.5 Automated testing", 2)
    doc.add_paragraph(
        "The latest saved test log reports 33 passed tests and 4 passed subtests in 19.09 seconds. "
        "This expanded test result supersedes the older 4/4 statement in the initial exercise notes. "
        "Passing tests cover code paths; they do not independently validate scientific accuracy."
    )
    log = REPORTS / "exercise5_test_results.txt"
    if log.exists():
        p = doc.add_paragraph()
        r = p.add_run(log.read_text(encoding="utf-8", errors="replace")[-1800:])
        r.font.name = "Consolas"; r.font.size = Pt(8)
    figure(doc, "Fig_10_pytest_results.png", "Figure 10. Evidence of the latest automated test run.")

    doc.add_heading("5. Additional Exploratory Indicators", 1)
    for x in [
        "Class-5 vegetation occupancy is a grid occupancy proxy, not a validated canopy-cover percentage.",
        "Slope bins are descriptive terrain categories, not a validated hazard prediction.",
        "Building-height estimates are sensitive to the local-ground search radius and are exploratory, not reliable height measurements.",
        "Impervious surface cannot be reliably estimated because ground class 2 can include paved and unpaved ground.",
        "Phase I/II land-cover datasets were not available as inputs, so verified cross-phase land-cover fusion is not claimed.",
    ]: bullet(doc, x)

    doc.add_heading("6. Limitations", 1)
    for x in [
        "Street imagery and LiDAR are not spatially matched at image-footprint level.",
        "Segmentation predictions have not been validated against manually labelled ground truth.",
        "No validated walkability or site-suitability model was produced.",
        "Vegetation occupancy is not canopy cover; available class labels do not support reliable impervious-surface estimation.",
        "Building-height estimates are exploratory and sensitive to the ground-neighbour search radius.",
        "The two profile boxes are within the same study area and are not proven representative or contrasting sites.",
        "Earlier notes contain slightly different DEM statistics from a prior raster summary. This report uses the latest stated GDAL values; differences can arise from statistics methods and NoData handling.",
        "If any figure is missing, the report records that it could not be inserted; check the local figure path.",
    ]: bullet(doc, x)

    doc.add_heading("7. Conclusion", 1)
    doc.add_paragraph(
        "Exercise 5 establishes a working, test-covered workflow for inspecting classified LiDAR, "
        "extracting ground points, deriving terrain rasters, summarizing metrics, creating JSON site "
        "profiles and incorporating existing street-scene segmentation outputs as exploratory context. "
        "The expanded test suite passed (33 tests and 4 subtests). Further work is needed before making "
        "scientific claims about canopy cover, impervious surfaces, building heights, walkability or site "
        "suitability. Recommended next steps are to verify all figures and parameters, validate segmentation "
        "against labelled samples, spatially align imagery and LiDAR where possible, and incorporate "
        "Phase I/II data only when available and documented."
    )
    doc.add_heading("Appendix A. Main output files", 1)
    for x in [
        "outputs/reports/unified_site_profile.json",
        "outputs/reports/unified_site_profile_enhanced.json",
        "outputs/reports/terrain_metrics.csv",
        "outputs/reports/exercise5_test_results.txt",
        "docs/exercise5/figures/Fig_01 through Fig_10",
    ]: bullet(doc, x)

    doc.save(OUTPUT)
    print("REPORT CREATED:", OUTPUT)
    print("REPORT SIZE (bytes):", OUTPUT.stat().st_size)
    print("FIGURE CHECK:")
    for i in range(1, 11):
        found = list(FIGURES.glob(f"Fig_{i:02d}_*"))
        print(f"  Figure {i:02d}:", "FOUND" if found else "MISSING")

if __name__ == "__main__":
    main()
